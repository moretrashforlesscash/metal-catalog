"""Три заявки от 29.10.2026 → три Excel-файла с ценами из каталога «Металл Сервис».

Цена без логистики: ₽ с НДС РФ 22% → без НДС → ₸ по курсу НБ РК → + НДС РК 16%.
Формулы пересчёта живут в Excel (курс и ставки НДС — ячейки вверху листа, их можно менять).
Запуск: python3 zayavki/build_2026-10-29.py
"""

import csv
import glob
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import find  # noqa: E402

RATE, RATE_DATE = 5.67, "29.10.2026"
ROWS = [r for f in sorted(glob.glob(str(ROOT / "catalog" / "metallservis*.csv"))) for r in csv.DictReader(open(f, encoding="utf-8"))]
PRICE_DATE = max(r["price_date"] for r in ROWS)


def pick(sec, name, size, wh=None):
    """Самая дешёвая нормальная (не некондиция) позиция: Москва, если есть, иначе любой склад. + минимум по всем складам."""
    c = [r for r in ROWS if re.search(sec, r["section"], re.I) and re.search(name, r["name"], re.I)
         and r["size"] == size and not find.is_offgrade(r) and r["unit"] in ("т", "теор.т")]
    if not c:
        return None, None
    cheapest = lambda rows: min(rows, key=lambda r: float(r["price_rub"]))  # noqa: E731
    msk = [r for r in c if r["warehouse"] == (wh or "Москва")]
    return cheapest(msk or c), cheapest(c)


def label(r):
    return f"{r['section']} | {r['name']} | {r['size_label']} {r['size']}".strip()


# (позиция из заявки, кол-во, ед., [(соответствие, раздел, наименование, размер, примечание, масса вручную)], цена клиента ₸/т)
NO = "нет в каталоге"
R1 = [
    ("Лист стальной г/к S2", None, "т", [("точно", "ЛИСТОВАЯ Г/К ОБЫЧ", "н/обр", "2", "Формат листа в прайсе не указан (обычно 1250×2500) — уточнить", None)], None),
    ("Лист стальной г/к S3", None, "т", [("точно", "ЛИСТОВАЯ Г/К ОБЫЧ", "н/обр", "3", "Формат листа в прайсе не указан — уточнить", None)], None),
    ("Лист стальной г/к S6", None, "т", [("точно", "ЛИСТОВАЯ Г/К ОБЫЧ", "н/обр", "6", "Формат листа в прайсе не указан — уточнить", None)], None),
    ("Арматура А500С Ø10", None, "т", [("точно", "^АРМАТУРА$", "А3 А400;500", "10", "А400/А500 в одной строке прайса", None)], 347700),
    ("Арматура А500С Ø12", None, "т", [("точно", "^АРМАТУРА$", "А3 А400;500", "12", "А400/А500 в одной строке прайса", None)], 364300),
    ("Арматура А500С Ø20", None, "т", [("точно", "^АРМАТУРА$", "А3 А400;500", "20", "А400/А500 в одной строке прайса", None)], 325500),
    ("Проволока вязальная", None, "т", [("аналог", "ПРОВОЛОКА О/К ТОРГ", "^отожженная", "1.2",
                                         "«Вязальная» в прайсе нет; это отожжённая проволока. Диаметр не указан — взят 1,2 мм (ходовой для вязки), есть 0,6–6 мм", None)], None),
    ("Двутавр 10", None, "т", [("точно", "БАЛКИ ДВУТАВР", "^10$", "", "Обычный двутавр ГОСТ 8239", None)], None),
    ("Двутавр 12", None, "т", [("аналог", "БАЛКИ ДВУТАВР", "^12$", "Б1", "Обычного 12 (ГОСТ 8239) нет — 12Б1 (ГОСТ 26020), легче: 8,7 против 11,5 кг/м", None),
                               ("аналог", "БАЛКИ ДВУТАВР", "^12$", "Б2", "", None)], None),
    ("Двутавр 16", None, "т", [("аналог", "БАЛК", "^16$", "Б1", "Обычного 16 нет ни на одном складе; 16Б1 есть только в Новосибирске", None)], None),
    ("Двутавр 20", None, "т", [("точно", "БАЛК", "^20$", "", "Обычный 20 есть только на Урале (в Москве нет)", None),
                               ("аналог", "БАЛКИ ДВУТАВР", "^20$", "Б1", "20Б1 есть в Москве", None)], None),
    ("Труба б/ш 57×3,5", None, "т", [("точно", "Г/Д", r"^57 Ст20", "3.5", "Горячедеформированная ГОСТ 8732, Ст20", None)], 668000),
    ("Труба б/ш 76×3,5", None, "т", [("точно", "Г/Д", r"^76 Ст20", "3.5", "Горячедеформированная ГОСТ 8732, Ст20", None)], 559000),
    ("Труба б/ш 76×4", None, "т", [("точно", "Г/Д", r"^76 Ст20", "4", "Та же строка прайса, что и 76×3,5", None)], None),
    ("Труба б/ш 108×6", None, "т", [("точно", "Г/Д", r"^108 Ст20", "6", "Горячедеформированная ГОСТ 8732, Ст20", None)], 559000),
    ("Труба б/ш 219×10", None, "т", [("точно", "Г/Д", r"^219 Ст20", "10", "Горячедеформированная ГОСТ 8732, Ст20", None)], 668000),
]
R2 = [
    ("Труба стальная Ø219*6 ГОСТ 10704-91", 24, "п/м", [("точно", "^ТРУБЫ ЭЛЕКТРОСВАРНЫЕ$", "^219", "6", "", None)], None),
    ("Сталь угловая ГОСТ 8509-93 50*50*4", 36, "п/м", [("точно", "^УГОЛОК$", "^50x50", "4", "", None)], None),
    ("Профильная труба ГОСТ 30245-2003/13663-86 100*100*3", 24, "п/м", [("точно", "ТРУБЫ ЭЛЕКТРОСВАРНЫЕ КВАДРАТ", r"^100\b", "3", "", None)], None),
    ("Труба стальная Ø76*3 ГОСТ 10704-91", 12, "п/м", [("точно", "^ТРУБЫ ЭЛЕКТРОСВАРНЫЕ$", "^76", "3", "", None)], None),
]
R3 = [
    ("Труба профильная 300*300 мм, стенка 3 мм", 48, "м", [("аналог", "ТРУБЫ ЭЛЕКТРОСВАРНЫЕ КВАДРАТ", r"^300\b", "6", "Стенки 3 мм для 300×300 нет — минимальная в каталоге 6 мм (вдвое тяжелее)", None)], None),
    ("Труба профильная 200*200 мм, стенка 3 мм", 60, "м", [("аналог", "ТРУБЫ ЭЛЕКТРОСВАРНЫЕ КВАДРАТ", r"^200\b", "5", "Стенки 3 мм для 200×200 нет — минимальная в каталоге 5 мм", None)], None),
    ("Труба профильная 150*150 мм, стенка 3 мм", 18, "м", [("аналог", "ТРУБЫ ЭЛЕКТРОСВАРНЫЕ КВАДРАТ", r"^150\b", "4", "Стенки 3 мм для 150×150 нет — минимальная в каталоге 4 мм", None)], None),
    ("Лист г/к рифлёный, толщина 4 мм", 42, "м2", [("точно", "ЛИСТ РИФЛЕНЫЙ", "чечев", "4", "Рифление «чечевица». Масса 31,4 кг/м² — как у гладкого листа; рифлёный тяжелее на несколько %, уточнить", 31.4),
                                                   ("аналог", "ЛИСТ РИФЛЕНЫЙ", "ромб", "4", "Рифление «ромб». Масса — как у гладкого листа, уточнить", 31.4)], None),
    ("Лист г/к, толщина 22 мм", 20, "м2", [("точно", "ЛИСТОВАЯ Г/К ОБЫЧ", "", "22", "Формат листа — уточнить", None)], None),
    ("Лист г/к, толщина 12 мм", 3, "м2", [("точно", "ЛИСТОВАЯ Г/К ОБЫЧ", "", "12", "Формат листа — уточнить", None)], None),
    ("Двутавровая балка, высота 300 мм", 7, "м", [("уточнить", "БАЛКИ ДВУТАВР", "^30$", "Б1", "Серия не указана: 30Б1 — самая лёгкая", None),
                                                  ("уточнить", "БАЛКИ ДВУТАВР", "^30$", "Ш1", "30Ш1 — широкополочная, тяжелее", None)], None),
    ("Угольник 40*40", 300, "м", [("уточнить", "^УГОЛОК$", "^40x40", "4", "Толщина не указана — взят 40×40×4", None),
                                  ("уточнить", "^УГОЛОК$", "^40x40", "3", "Вариант 40×40×3", None)], None),
    ("Арматура диам 16", 90, "м", [("точно", "^АРМАТУРА$", "А3 А400;500", "16", "Класс не указан — взята А500С (А3)", None)], None),
    ("Швеллер горячекатанный 300*60*5", 24, "м", [("аналог", "^ШВЕЛЛЕР$", "^30$", "", "Г/к швеллера 300×60 не бывает; №30 (ГОСТ 8240) — 300×100", None),
                                                  ("аналог", "ШВЕЛЛЕР ГНУТЫЙ", "^300x100", "6", "Гнутого 300×60×5 нет; ближайший гнутый 300×100×6", None)], None),
    ("Лебёдка с редуктором 500 кг, канат 30 м, 380 В, TOR", 1, "шт", [], None),
    ("Блок монтажный опорный TOR 5 т", 1, "шт", [], None),
    ("Труба диам 400 мм, стенка 8–10 мм", 24, "шт", [("аналог", "^ТРУБЫ ЭЛЕКТРОСВАРНЫЕ$", "^426", "8", "Ø400 нет; стандартный размер — 426. Длина трубы не указана → сумма не посчитана", None),
                                                    ("аналог", "^ТРУБЫ ЭЛЕКТРОСВАРНЫЕ$", "^426", "10", "Вариант со стенкой 10 мм", None)], None),
]

ARIAL = "Arial"
BLUE, BLACK = Font(name=ARIAL, color="0000FF"), Font(name=ARIAL)
BOLD = Font(name=ARIAL, bold=True)
HEAD = Font(name=ARIAL, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="305496")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY, KG = '#,##0;-#,##0;"-"', "0.000"


def build(path, title, items, client_col):
    wb = Workbook()
    ws = wb.active
    ws.title = "Расчёт"
    ws["A1"] = title
    ws["A1"].font = Font(name=ARIAL, bold=True, size=14)
    params = [("Курс ₸ за 1 ₽", RATE, f"Официальный курс НБ РК на {RATE_DATE} (scripts/rate.py). Реальный курс оплаты может отличаться", "0.00"),
              ("НДС РФ в цене поставщика", 0.22, "Цены «Металл Сервис» — с НДС РФ 22% (с 01.01.2026)", "0%"),
              ("НДС РК", 0.16, "Налоговый кодекс РК, с 01.01.2026", "0%")]
    for i, (k, v, note, fmt) in enumerate(params, start=3):
        ws.cell(i, 1, k).font = BLACK
        c = ws.cell(i, 2, v)
        c.font, c.fill, c.number_format = BLUE, YELLOW, fmt
        ws.cell(i, 3, note).font = Font(name=ARIAL, italic=True, color="808080")
    ws["A7"] = (f"Цены из прайса «Металл Сервис» от {PRICE_DATE.split('-')[2]}.{PRICE_DATE.split('-')[1]}.{PRICE_DATE[:4]}, без доставки. "
                "₸/т с НДС РК = ₽/т ÷ (1 + НДС РФ) × курс × (1 + НДС РК). Вес теоретический (факт ±3–8%). "
                "Синие ячейки на жёлтом — можно менять, остальное пересчитается.")
    ws["A7"].font = Font(name=ARIAL, italic=True)
    ws["A7"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A7:L7")
    ws.row_dimensions[7].height = 32

    cols = [("№", "№"), ("text", "Позиция из заявки"), ("qty", "Кол-во"), ("unit", "Ед."), ("found", "Что нашли в каталоге"),
            ("status", "Соответствие"), ("wh", "Склад"), ("rub", "Цена, ₽/т\n(с НДС РФ 22%)"), ("mass", "Масса,\nкг/м (кг/м²)"),
            ("kzt", "Цена, ₸/т\nс НДС 16%"), ("kzt_net", "Цена, ₸/т\nбез НДС"), ("kzt_m", "Цена, ₸ за м (м²)\nс НДС 16%"),
            ("sum", "Сумма, ₸\nс НДС 16%")]
    if client_col:
        cols += [("client", "Цена в заявке,\n₸/т"), ("diff", "Наша цена\nк цене заявки")]
    cols += [("best_wh", "Дешевле всего:\nсклад"), ("best_rub", "₽/т на том\nскладе"), ("best_kzt", "₸/т с НДС 16%\nна том складе"),
             ("src", "Источник\n(файл:строка)"), ("note", "Примечание")]
    H = 9
    for j, (_, name) in enumerate(cols, start=1):
        c = ws.cell(H, j, name)
        c.font, c.fill, c.border = HEAD, HEAD_FILL, BOX
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[H].height = 45
    col = {k: j for j, (k, _) in enumerate(cols, start=1)}
    L = {k: ws.cell(1, j).column_letter for k, j in col.items()}
    tn = lambda ref: f"={ref}/(1+$B$4)*$B$3*(1+$B$5)"  # noqa: E731

    r, n, sum_cells = H + 1, 0, []
    for text, qty, unit, cands, client in items:
        n += 1
        found = [(c, pick(c[1], c[2], c[3])) for c in cands]
        found = [(c, p) for c, p in found if p[0]]
        lines = found or [(None, (None, None))]
        for k, (cand, (row, best)) in enumerate(lines):
            first = k == 0
            vals = {"№": n if first else None, "text": text if first else "↳ вариант",
                    "qty": qty if first else None, "unit": unit if first else None}
            for name, v in vals.items():
                ws.cell(r, col[name], v).font = BLACK
            if row is None:
                ws.cell(r, col["status"], NO).font = Font(name=ARIAL, bold=True, color="C00000")
                ws.cell(r, col["note"], "Не металлопрокат — в прайсе «Металл Сервис» нет").font = BLACK
            else:
                status, *_, note, mass_manual = cand
                mass = mass_manual or find.mass_kg_m(row)
                ws.cell(r, col["found"], label(row)).font = BLACK
                ws.cell(r, col["status"], status).font = Font(name=ARIAL, bold=True,
                                                                    color={"точно": "007A33", "аналог": "C55A11"}.get(status, "7F6000"))
                ws.cell(r, col["wh"], row["warehouse"]).font = BLACK
                c = ws.cell(r, col["rub"], float(row["price_rub"]))
                c.font, c.number_format = BLUE, MONEY
                c.comment = Comment(f"{row['unit']} — {'теоретическая тонна' if row['unit'] == 'теор.т' else 'тонна'}; "
                                    f"прайс {row['price_date']}", "metal-catalog")
                if mass:
                    c = ws.cell(r, col["mass"], round(mass, 3))
                    c.font, c.number_format = (BLUE if mass_manual else BLACK), KG
                ws.cell(r, col["kzt"], tn(f"{L['rub']}{r}")).number_format = MONEY
                ws.cell(r, col["kzt_net"], f"={L['kzt']}{r}/(1+$B$5)").number_format = MONEY
                if mass:
                    ws.cell(r, col["kzt_m"], f"={L['kzt']}{r}*{L['mass']}{r}/1000").number_format = MONEY
                    if first and qty and unit in ("м", "п/м", "м2"):
                        ws.cell(r, col["sum"], f"=C{r}*{L['kzt_m']}{r}").number_format = MONEY
                        sum_cells.append(f"{L['sum']}{r}")
                if client_col and client and first:
                    c = ws.cell(r, col["client"], client)
                    c.font, c.number_format = BLUE, MONEY
                    ws.cell(r, col["diff"], f"={L['kzt']}{r}/{L['client']}{r}-1").number_format = '+0%;-0%;0%'
                ws.cell(r, col["best_wh"], best["warehouse"]).font = BLACK
                c = ws.cell(r, col["best_rub"], float(best["price_rub"]))
                c.font, c.number_format = BLUE, MONEY
                ws.cell(r, col["best_kzt"], tn(f"{L['best_rub']}{r}")).number_format = MONEY
                ws.cell(r, col["src"], f"{row['src']} · {row['price_date']}").font = BLACK
                extra = " Цена за теоретический вес." if row["unit"] == "теор.т" else ""
                ws.cell(r, col["note"], (note + extra).strip()).font = BLACK
            for j in range(1, len(cols) + 1):
                cell = ws.cell(r, j)
                cell.border = BOX
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.font == Font():
                    cell.font = BLACK
                if not first:
                    cell.fill = GREY
            r += 1
    if sum_cells:
        ws.cell(r, col["sum"] - 1, "Итого (по первым вариантам):").font = BOLD
        c = ws.cell(r, col["sum"], "=" + "+".join(sum_cells))
        c.font, c.number_format = BOLD, MONEY
        r += 1
    r += 1
    ws.cell(r, 1, "Соответствие: «точно» — та же позиция; «аналог» — точной нет, предложена ближайшая; "
                  "«уточнить» — в заявке не хватает данных, взят наиболее вероятный вариант. Серые строки — варианты.").font = Font(name=ARIAL, italic=True)
    r += 1
    ws.cell(r, 1, "«Дешевле всего» — минимум по складам Москва/Екатеринбург/Челябинск/Пермь/Новосибирск/Омск без учёта доставки.").font = Font(name=ARIAL, italic=True)

    widths = {"text": 34, "qty": 7, "unit": 5, "found": 42, "status": 12, "wh": 13, "src": 26, "note": 48}
    for name, j in col.items():
        ws.column_dimensions[ws.cell(1, j).column_letter].width = widths.get(name, 14)
    ws.column_dimensions["A"].width = 26  # параметры вверху
    ws.freeze_panes = ws.cell(H + 1, 3)
    for row_cells in ws.iter_rows():
        for cell in row_cells:
            if cell.value is not None and cell.font.name != ARIAL:
                cell.font = Font(name=ARIAL, bold=cell.font.bold, color=cell.font.color)
    wb.save(path)


if __name__ == "__main__":
    out = ROOT / "zayavki"
    build(out / "2026-10-29_Заявка_1.xlsx", "Заявка 1 — цены без логистики", R1, client_col=True)
    build(out / "2026-10-29_Заявка_2.xlsx", "Заявка 2 (энергоснабжение, ЯКНО) — цены без логистики", R2, client_col=False)
    build(out / "2026-10-29_Заявка_3.xlsx", "Заявка 3 (ТК, рама установки) — цены без логистики", R3, client_col=False)
    print("готово:", *sorted(p.name for p in out.glob("2026-10-29_*.xlsx")))
